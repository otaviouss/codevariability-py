"""Regression checks for untrusted labels in exported spreadsheets and paths."""

import csv
import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from codevariability import analyze, compare_groups


class ExportSecurityTest(unittest.TestCase):
    def test_external_metric_cannot_write_outside_export_directory(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "sample.py"
            source.write_text("x = 1", encoding="utf-8")
            result = analyze([source], metrics="cosine")
            external = pd.DataFrame([[1.0]], index=result.files, columns=result.files)
            result = result.with_metric("../escaped", external)
            output = root / "export"
            result.export(output, formats="csv")
            self.assertEqual(len(list(output.glob("*similarity_matrix.csv"))), 2)
            self.assertFalse(list(root.glob("escaped*")))

    def test_untrusted_file_and_group_names_are_plain_spreadsheet_text(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            suspicious_names = ("=1+1.py", "+1.py", "-1.py", "@SUM(1,1).py")
            sources = []
            for suspicious in suspicious_names:
                source = root / suspicious
                source.write_text("x = 1", encoding="utf-8")
                sources.append(source)
            other = root / "normal.py"
            other.write_text("x = 2", encoding="utf-8")
            result = analyze([*sources, other], metrics="cosine")
            output = root / "export"
            result.export(output)
            result.to_excel(output / "analysis.xlsx")

            with (output / "cosine_similarity_matrix.csv").open(newline="", encoding="utf-8") as handle:
                rows = list(csv.reader(handle))
            for suspicious in suspicious_names:
                self.assertIn("\t" + suspicious, rows[0])
                self.assertIn("\t" + suspicious, [row[0] for row in rows[1:]])
                self.assertIn(suspicious, result.files)
            workbook = load_workbook(output / "analysis.xlsx")
            for sheet in workbook:
                for row in sheet:
                    for cell in row:
                        self.assertNotEqual(cell.data_type, "f")

    def test_multigroup_export_escapes_index_labels(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            groups = {}
            for name in ("=1+1", "regular", "@SUM(1,1)"):
                folder = root / name
                folder.mkdir()
                for index in range(2):
                    (folder / f"{index}.py").write_text(f"x = {index}", encoding="utf-8")
                groups[name] = folder
            result = compare_groups(groups, metrics="cosine", permutations=2)
            output = root / "export"
            result.export(output)
            result.to_excel(output / "groups.xlsx")
            with (output / "pairwise.csv").open(newline="", encoding="utf-8") as handle:
                rows = list(csv.reader(handle))
            self.assertTrue(any("\t=1+1" in row for row in rows))
            workbook = load_workbook(output / "groups.xlsx")
            for sheet in workbook:
                for row in sheet:
                    for cell in row:
                        self.assertNotEqual(cell.data_type, "f")


if __name__ == "__main__":
    unittest.main()
